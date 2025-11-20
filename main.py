"""Utility to convert XLSX-based test case tables into formatted text blocks."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Dict, Iterator, List, Optional

from openpyxl import load_workbook


# Authoritative column order inside each table.
EXPECTED_HEADERS: List[str] = [
	"test case id",
	"requirements",
	"level",
	"description",
	"pre-conditions",
	"test steps",
	"test data",
	"expected result",
	"actual result",
	"status",
]


def normalize(value: Optional[str]) -> str:
	return str(value).strip().lower() if value is not None else ""


def is_header_row(row: List[Optional[str]]) -> bool:
	normalized = [normalize(cell) for cell in row]
	return normalized == EXPECTED_HEADERS or normalized == [
		header.replace("description", "desciption") for header in EXPECTED_HEADERS
	]


def iter_test_rows(sheet) -> Iterator[Dict[str, Optional[str]]]:  # type: ignore[no-untyped-def]
	header_found = False
	headers: List[str] = []

	for row in sheet.iter_rows(values_only=True):
		values = list(row)

		if not any(values):
			header_found = False
			headers = []
			continue

		if is_header_row(values):
			header_found = True
			# Use canonical header names to simplify downstream logic.
			headers = EXPECTED_HEADERS.copy()
			continue

		if not header_found:
			continue

		record = {
			header: (value if value is not None else "")
			for header, value in zip(headers, values)
		}
		if record.get("test case id"):
			yield record


def format_case(record: Dict[str, Optional[str]]) -> str:
	def get(field: str) -> str:
		value = record.get(field, "")
		return str(value).strip() if value is not None else ""

	return "\n".join(
		[
			f"Test Case ID: <{get('test case id') or 'N/A'}>",
			f"Requirements: <{get('requirements') or 'N/A'}>",
			f"Level: <{get('level') or 'N/A'}>",
			f"Description: <{get('description') or 'N/A'}>",
			f"Pre-Conditions: <{get('pre-conditions') or 'N/A'}>",
			f"Test Steps: <{get('test steps') or 'N/A'}>",
			f"Test Data: <{get('test data') or 'N/A'}>",
			f"Expected Result: <{get('expected result') or 'N/A'}>",
		]
	)


def parse_workbook(path: Path, sheet_name: Optional[str]) -> List[str]:
	workbook = load_workbook(path, data_only=True)
	sheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets

	formatted_cases: List[str] = []
	for sheet in sheets:
		for record in iter_test_rows(sheet):
			formatted_cases.append(format_case(record))

	return formatted_cases


def build_arg_parser() -> argparse.ArgumentParser:
	parser = argparse.ArgumentParser(
		description="Convert XLSX test case tables into formatted text output."
	)
	parser.add_argument("xlsx", type=Path, help="Path to the XLSX file to parse")
	parser.add_argument(
		"--sheet",
		help="Optional sheet name to parse; parses all sheets when omitted",
	)
	return parser


def main() -> None:
	parser = build_arg_parser()
	args = parser.parse_args()

	if not args.xlsx.exists():
		raise FileNotFoundError(f"Could not find XLSX file at {args.xlsx}")

	cases = parse_workbook(args.xlsx, args.sheet)

	if not cases:
		print("No test cases found with the expected headers.")
		return

	print("\n\n".join(cases))


if __name__ == "__main__":
	main()
